"""Local data and pricing rules for Papéis de Marte. No network services."""
from __future__ import annotations

import os
import math
import sqlite3
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path


def money(value: int | None) -> str:
    if value is None:
        return "A definir"
    value = int(value)
    sign = "-" if value < 0 else ""
    whole, cents = divmod(abs(value), 100)
    return f"{sign}R$ {whole:,}".replace(",", ".") + f",{cents:02d}"


def cents(text: str, *, allow_empty: bool = False) -> int | None:
    raw = str(text).strip().replace("R$", "").replace(" ", "")
    if not raw and allow_empty:
        return None
    if "," in raw:
        raw = raw.replace(".", "").replace(",", ".")
    try:
        from decimal import Decimal, ROUND_HALF_UP
        value = Decimal(raw)
        if not value.is_finite() or value < 0:
            raise ValueError()
        return int((value * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except Exception as exc:
        raise ValueError("Informe um valor válido, como 12,50.") from exc


def quantity(text: str) -> float:
    try:
        value = float(str(text).strip().replace(",", "."))
        if not 0 < value <= 1000000:
            raise ValueError()
        return value
    except Exception as exc:
        raise ValueError("A quantidade precisa ser maior que zero.") from exc


def fmt_qty(value: float) -> str:
    return f"{value:g}".replace(".", ",")


def parse_date(text: str) -> str:
    try:
        return datetime.strptime(text.strip(), "%d/%m/%Y").date().isoformat()
    except ValueError as exc:
        raise ValueError("Informe a data de entrega em DD/MM/AAAA.") from exc


def br_date(value: str | None) -> str:
    return date.fromisoformat(value).strftime("%d/%m/%Y") if value else "Não informada"


def automatic_code(name: str, size: str = '', grammage: str = '') -> str:
    """Three letters of name + normalized size + grammage, e.g. OFFA4150."""
    plain = unicodedata.normalize('NFKD', name.strip()).encode('ascii', 'ignore').decode('ascii')
    prefix = ''.join(ch for ch in plain.upper() if ch.isalpha())[:3]
    if len(prefix) < 3:
        raise ValueError('O nome deve conter ao menos três letras para gerar o código.')
    def part(value):
        normalized = unicodedata.normalize('NFKD', str(value).strip()).encode('ascii', 'ignore').decode('ascii')
        return ''.join(ch for ch in normalized.upper() if ch.isalnum() or ch == '.')
    return prefix + part(size) + part(grammage)


def asset(name: str) -> Path:
    return Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)) / "assets" / name


def data_directory() -> Path:
    if sys.platform == "win32":
        base = Path(os.environ.get("LOCALAPPDATA", str(Path.home() / "AppData" / "Local")))
    else:
        base = Path(os.environ.get("XDG_DATA_HOME", str(Path.home() / ".local" / "share")))
    path = base / "ERP_Papeis_de_Marte"
    path.mkdir(parents=True, exist_ok=True)
    return path


def filter_orders(orders, filters):
    """Apply independent column filters to the displayed order records."""
    mapping={'number':'number','customer':'customer','items':'descriptions',
             'created':'created_date','due':'due_date','payment':'payment','production':'production'}
    filtered=[]
    for order in orders:
        if any((filters.get(key) or '').casefold() not in
               (br_date(order[field]) if key in ('created','due') else str(order[field] or '')).casefold()
               for key,field in mapping.items()):
            continue
        total=None if order['unpriced'] else order['total_cents']
        if filters.get('min_total') is not None and (total is None or total < filters['min_total']):
            continue
        if filters.get('max_total') is not None and (total is None or total > filters['max_total']):
            continue
        remaining=order['remaining_cents']
        if filters.get('min_remaining') is not None and (remaining is None or remaining < filters['min_remaining']):
            continue
        if filters.get('max_remaining') is not None and (remaining is None or remaining > filters['max_remaining']):
            continue
        filtered.append(order)
    return filtered


class Store:
    def __init__(self, path: str | Path, initial_workbook: str | Path | None = None):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(self.path)
        self.db.row_factory = sqlite3.Row
        self.db.execute("PRAGMA foreign_keys = ON")
        self.db.executescript("""
          CREATE TABLE IF NOT EXISTS materials (
            id INTEGER PRIMARY KEY, code TEXT NOT NULL UNIQUE COLLATE NOCASE,
            name TEXT NOT NULL, specification TEXT NOT NULL DEFAULT '',
            size TEXT NOT NULL DEFAULT '', grammage TEXT NOT NULL DEFAULT '',
            pack_qty REAL NOT NULL CHECK(pack_qty > 0), unit TEXT NOT NULL DEFAULT 'un',
            pack_cents INTEGER NOT NULL CHECK(pack_cents >= 0), active INTEGER NOT NULL DEFAULT 1
          );
          CREATE TABLE IF NOT EXISTS material_variants (
            id INTEGER PRIMARY KEY, material_id INTEGER NOT NULL REFERENCES materials(id),
            name TEXT NOT NULL COLLATE NOCASE, active INTEGER NOT NULL DEFAULT 1,
            UNIQUE(material_id,name)
          );
          CREATE TABLE IF NOT EXISTS products (
            id INTEGER PRIMARY KEY, code TEXT NOT NULL UNIQUE COLLATE NOCASE,
            name TEXT NOT NULL, size TEXT NOT NULL DEFAULT '', grammage TEXT NOT NULL DEFAULT '',
            markup REAL NOT NULL DEFAULT 1.8 CHECK(markup >= 1),
            table_cents INTEGER CHECK(table_cents >= 0), active INTEGER NOT NULL DEFAULT 1
          );
          CREATE TABLE IF NOT EXISTS recipes (
            product_id INTEGER NOT NULL REFERENCES products(id),
            position INTEGER NOT NULL, material_code TEXT NOT NULL,
            qty REAL NOT NULL CHECK(qty > 0), PRIMARY KEY(product_id,position)
          );
          CREATE TABLE IF NOT EXISTS orders (
            id INTEGER PRIMARY KEY, number TEXT NOT NULL UNIQUE COLLATE NOCASE,
            customer TEXT NOT NULL, created_date TEXT, due_date TEXT,
            payment TEXT NOT NULL DEFAULT 'Pendente', paid_cents INTEGER NOT NULL DEFAULT 0,
            production TEXT NOT NULL DEFAULT 'Novo',
            notes TEXT NOT NULL DEFAULT ''
          );
          CREATE TABLE IF NOT EXISTS order_sequence (
            singleton INTEGER PRIMARY KEY CHECK(singleton=1), last_number INTEGER NOT NULL
          );
          CREATE TABLE IF NOT EXISTS order_items (
            id INTEGER PRIMARY KEY, order_id INTEGER NOT NULL REFERENCES orders(id) ON DELETE CASCADE,
            product_id INTEGER REFERENCES products(id) ON DELETE SET NULL,
            description TEXT NOT NULL, qty REAL NOT NULL CHECK(qty > 0),
            unit_cents INTEGER CHECK(unit_cents >= 0)
          );
          CREATE TABLE IF NOT EXISTS order_item_variants (
            item_id INTEGER NOT NULL REFERENCES order_items(id) ON DELETE CASCADE,
            material_id INTEGER NOT NULL REFERENCES materials(id),
            variant_id INTEGER NOT NULL REFERENCES material_variants(id),
            PRIMARY KEY(item_id,material_id)
          );
          CREATE TABLE IF NOT EXISTS stock_movements (
            id INTEGER PRIMARY KEY, material_id INTEGER NOT NULL REFERENCES materials(id),
            variant_id INTEGER REFERENCES material_variants(id),
            order_id INTEGER REFERENCES orders(id), delta REAL NOT NULL,
            kind TEXT NOT NULL, notes TEXT NOT NULL DEFAULT '',
            effective_date TEXT NOT NULL, created_at TEXT NOT NULL
          );
          CREATE INDEX IF NOT EXISTS stock_material_date ON stock_movements(material_id,effective_date);
          CREATE INDEX IF NOT EXISTS stock_order_id ON stock_movements(order_id);
        """)
        self.migrate_schema()
        if initial_workbook is not None and not self.db.execute("SELECT 1 FROM materials LIMIT 1").fetchone():
            self.import_workbook(initial_workbook)

    def migrate_schema(self):
        """Add fields without resetting a database already in use."""
        with self.db:
            order_columns={r['name'] for r in self.all('PRAGMA table_info(orders)')}
            if 'paid_cents' not in order_columns:
                self.db.execute('ALTER TABLE orders ADD COLUMN paid_cents INTEGER NOT NULL DEFAULT 0')
            material_columns={r['name'] for r in self.all('PRAGMA table_info(materials)')}
            product_columns={r['name'] for r in self.all('PRAGMA table_info(products)')}
            for name in ('size','grammage'):
                if name not in material_columns:
                    self.db.execute(f"ALTER TABLE materials ADD COLUMN {name} TEXT NOT NULL DEFAULT ''")
                if name not in product_columns:
                    self.db.execute(f"ALTER TABLE products ADD COLUMN {name} TEXT NOT NULL DEFAULT ''")
            stock_columns={r['name'] for r in self.all('PRAGMA table_info(stock_movements)')}
            if 'variant_id' not in stock_columns:
                self.db.execute('ALTER TABLE stock_movements ADD COLUMN variant_id INTEGER REFERENCES material_variants(id)')
            # Older builds kept these fields together in 'specification'.
            if 'size' not in material_columns or 'grammage' not in material_columns:
                for row in self.all("SELECT id,specification FROM materials WHERE specification<>''"):
                    parts=[p.strip() for p in row['specification'].split(' · ',1)]
                    self.db.execute('UPDATE materials SET size=?,grammage=?,specification=? WHERE id=?',
                                    (parts[0],parts[1] if len(parts)>1 else '', '',row['id']))

    def close(self):
        self.db.close()

    def all(self, query, params=()):
        return self.db.execute(query, params).fetchall()

    def one(self, query, params=()):
        return self.db.execute(query, params).fetchone()

    def materials(self):
        return self.all("SELECT * FROM materials ORDER BY active DESC, name COLLATE NOCASE, code")

    def variants(self, material_id, active_only=False):
        return self.all('SELECT * FROM material_variants WHERE material_id=?'+
                        (' AND active=1' if active_only else '')+' ORDER BY active DESC,name COLLATE NOCASE',
                        (material_id,))

    def variant_requirements(self, product_id):
        """One choice for each material in a recipe that has active variations."""
        result=[];seen=set()
        for part in self.recipe(product_id):
            m=self.one('SELECT id,name,code FROM materials WHERE code=? COLLATE NOCASE',(part['material_code'],))
            if m and m['id'] not in seen:
                options=self.variants(m['id'],active_only=True)
                if options:result.append((m,options))
                seen.add(m['id'])
        return result

    def variant_labels(self, selections):
        labels=[]
        for material_id,variant_id in sorted((int(k),int(v)) for k,v in (selections or {}).items()):
            row=self.one('''SELECT m.name AS material,v.name AS variant FROM material_variants v
                            JOIN materials m ON m.id=v.material_id
                            WHERE v.id=? AND m.id=?''',(variant_id,material_id))
            if row:labels.append(f"{row['material']}: {row['variant']}")
        return labels

    def save_material(self, *, id=None, name, size='', grammage='', specification='', pack_qty, unit, pack_cents, code=None, variants=None):
        name,size,grammage = name.strip(),size.strip(),grammage.strip()
        code = code.strip().upper() if code is not None else automatic_code(name,size,grammage)
        if not code or not name:
            raise ValueError("Preencha o código e o nome do insumo.")
        if pack_qty <= 0 or pack_cents < 0:
            raise ValueError("Quantidade e preço da embalagem inválidos.")
        if variants is not None:
            cleaned=[v.strip() for v in variants if v.strip()]
            if len({v.casefold() for v in cleaned}) != len(cleaned):
                raise ValueError('Não repita nomes de variações do mesmo insumo.')
        with self.db:
            if id:
                old=self.one('SELECT code FROM materials WHERE id=?',(id,))
                if old is None:
                    raise ValueError('Insumo não encontrado.')
                self.db.execute("UPDATE materials SET code=?,name=?,size=?,grammage=?,specification=?,pack_qty=?,unit=?,pack_cents=? WHERE id=?",
                                (code,name,size,grammage,specification.strip(),pack_qty,unit.strip() or 'un',pack_cents,id))
                if old['code'].upper()!=code:
                    self.db.execute('UPDATE recipes SET material_code=? WHERE material_code=? COLLATE NOCASE',
                                    (code,old['code']))
            else:
                id=self.db.execute("INSERT INTO materials(code,name,size,grammage,specification,pack_qty,unit,pack_cents) VALUES(?,?,?,?,?,?,?,?)",
                                   (code,name,size,grammage,specification.strip(),pack_qty,unit.strip() or 'un',pack_cents)).lastrowid
            if variants is not None:
                names={v.casefold() for v in cleaned}
                for previous in self.variants(id):
                    self.db.execute('UPDATE material_variants SET active=? WHERE id=?',
                                    (int(previous['name'].casefold() in names),previous['id']))
                for value in cleaned:
                    self.db.execute('''INSERT INTO material_variants(material_id,name,active) VALUES(?,?,1)
                                       ON CONFLICT(material_id,name) DO UPDATE SET active=1''',(id,value))

    def toggle_material(self, id):
        with self.db:
            self.db.execute("UPDATE materials SET active=1-active WHERE id=?", (id,))

    def products(self):
        return self.all("SELECT * FROM products ORDER BY active DESC, name COLLATE NOCASE")

    def recipe(self, product_id):
        return self.all("""SELECT r.*, m.name AS material_name, m.pack_qty, m.pack_cents, m.unit
           FROM recipes r LEFT JOIN materials m ON m.code=r.material_code COLLATE NOCASE
           WHERE r.product_id=? ORDER BY r.position""", (product_id,))

    def product_cost(self, product_id):
        rows = self.recipe(product_id)
        missing = [r['material_code'] for r in rows if r['material_name'] is None]
        if not rows or missing:
            return None, missing
        return sum(r['qty'] * r['pack_cents'] / r['pack_qty'] for r in rows), []

    def suggested(self, product):
        cost, missing = self.product_cost(product['id'])
        if cost is None:
            return None, missing
        from decimal import Decimal, ROUND_HALF_UP
        price = int((Decimal(str(cost)) * Decimal(str(product['markup']))).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
        return price, []

    def save_product(self, *, id=None, name, size='', grammage='', markup, table_cents, recipe, code=None):
        name,size,grammage = name.strip(),size.strip(),grammage.strip()
        code = code.strip().upper() if code is not None else automatic_code(name,size,grammage)
        if not code or not name or not math.isfinite(markup) or markup < 1:
            raise ValueError("Preencha código e nome; o multiplicador deve ser no mínimo 1.")
        if table_cents is not None and table_cents < 0:
            raise ValueError("Preço de tabela inválido.")
        if not recipe:
            raise ValueError("Adicione ao menos um insumo à composição.")
        with self.db:
            if id:
                self.db.execute("UPDATE products SET code=?,name=?,size=?,grammage=?,markup=?,table_cents=? WHERE id=?",
                                (code,name,size,grammage,markup,table_cents,id))
                self.db.execute("DELETE FROM recipes WHERE product_id=?", (id,))
            else:
                id = self.db.execute("INSERT INTO products(code,name,size,grammage,markup,table_cents) VALUES(?,?,?,?,?,?)",
                                     (code,name,size,grammage,markup,table_cents)).lastrowid
            for pos,(material_code,qty) in enumerate(recipe):
                if not self.one("SELECT 1 FROM materials WHERE code=? COLLATE NOCASE", (material_code,)):
                    raise ValueError(f"Insumo não encontrado: {material_code}")
                self.db.execute("INSERT INTO recipes VALUES(?,?,?,?)", (id,pos,material_code,qty))
        return id

    def toggle_product(self, id):
        with self.db:
            self.db.execute("UPDATE products SET active=1-active WHERE id=?", (id,))

    def order(self, id):
        return self.one("SELECT * FROM orders WHERE id=?", (id,))

    def items(self, id):
        items=[]
        for row in self.all("SELECT * FROM order_items WHERE order_id=? ORDER BY id", (id,)):
            item=dict(row)
            item['variants']={v['material_id']:v['variant_id'] for v in self.all(
                'SELECT material_id,variant_id FROM order_item_variants WHERE item_id=?',(row['id'],))}
            items.append(item)
        return items

    def orders(self):
        return self.all("""SELECT o.*, SUM(CASE WHEN i.unit_cents IS NOT NULL THEN ROUND(i.qty*i.unit_cents) ELSE 0 END) AS total_cents,
          CASE WHEN SUM(CASE WHEN i.unit_cents IS NULL THEN 1 ELSE 0 END)>0 THEN NULL
               WHEN o.payment IN ('Pago','Presente') THEN 0
               ELSE CAST(MAX(0,SUM(ROUND(i.qty*i.unit_cents))-o.paid_cents) AS INTEGER) END AS remaining_cents,
          SUM(CASE WHEN i.unit_cents IS NULL THEN 1 ELSE 0 END) AS unpriced,
          GROUP_CONCAT(i.description || COALESCE((SELECT ' [' || GROUP_CONCAT(m.name||': '||v.name, ', ') || ']'
            FROM order_item_variants iv JOIN materials m ON m.id=iv.material_id
            JOIN material_variants v ON v.id=iv.variant_id WHERE iv.item_id=i.id),''), ', ') AS descriptions
          FROM orders o JOIN order_items i ON i.order_id=o.id
          GROUP BY o.id ORDER BY CASE WHEN o.due_date IS NULL THEN 1 ELSE 0 END, o.due_date, o.id DESC""")

    def next_number(self):
        numbers = [int(r[0]) for r in self.all("SELECT SUBSTR(number,3) FROM orders WHERE number GLOB 'PM[0-9][0-9][0-9][0-9]*'") if str(r[0]).isdigit()]
        saved=self.one('SELECT last_number FROM order_sequence WHERE singleton=1')
        return f"PM{max(max(numbers, default=0),saved['last_number'] if saved else 0)+1:04d}"

    def _remember_number(self, number):
        self.db.execute('''INSERT INTO order_sequence(singleton,last_number) VALUES(1,?)
                           ON CONFLICT(singleton) DO UPDATE SET last_number=MAX(last_number,excluded.last_number)''',
                        (int(number[2:]),))

    def delete_order(self, id):
        """Remove o pedido e devolve seus insumos, preservando o histórico de estoque."""
        with self.db:
            order=self.order(id)
            if order is None:
                raise ValueError('Pedido não encontrado.')
            last_number=int(self.next_number()[2:])-1
            self._remember_number(f'PM{last_number:04d}')
            movements=self.all('''SELECT material_id,variant_id,SUM(delta) AS net FROM stock_movements
                                  WHERE order_id=? GROUP BY material_id,variant_id''',(id,))
            for row in movements:
                if abs(row['net'])>1e-8:
                    self._record_movement(row['material_id'],-row['net'],'order_reversal',
                                          f"Estorno por exclusão do pedido {order['number']}",
                                          variant_id=row['variant_id'])
            self.db.execute('UPDATE stock_movements SET order_id=NULL WHERE order_id=?',(id,))
            self.db.execute('DELETE FROM order_items WHERE order_id=?',(id,))
            self.db.execute('DELETE FROM orders WHERE id=?',(id,))

    def save_order(self, *, id=None, customer, due_date, payment, production, notes, items, paid_cents=0):
        customer = customer.strip()
        if not customer or not items or not due_date:
            raise ValueError("Preencha cliente, entrega e ao menos um item.")
        date.fromisoformat(due_date)
        for item in items:
            if not item['description'].strip() or item['qty'] <= 0 or item['unit_cents'] is None or item['unit_cents'] < 0:
                raise ValueError("Cada item precisa de descrição, quantidade e preço.")
        if payment not in ('Pendente','Parcial','Pago','Presente'):
            raise ValueError('Selecione uma situação de pagamento válida.')
        if not isinstance(paid_cents,int) or paid_cents<0:
            raise ValueError('O valor recebido precisa ser válido e positivo.')
        total_cents=sum(round(item['qty']*item['unit_cents']) for item in items)
        if payment=='Parcial':
            if not 0<paid_cents<total_cents:
                raise ValueError('No pagamento parcial, o valor recebido deve ser maior que zero e menor que o total.')
        elif payment=='Pago':
            paid_cents=total_cents
        else:
            paid_cents=0
        old_items=list(self.items(id)) if id else []
        def signature(rows):
            return Counter((r['product_id'],round(float(r['qty']),8),
                            tuple(sorted((int(k),int(v)) for k,v in (r.get('variants') or {}).items()))) for r in rows)
        stock_change = not id or signature(old_items)!=signature(items)
        usage=defaultdict(float)
        if stock_change:
            for item in items:
                if item['product_id'] is None:
                    continue
                product=self.one('SELECT code FROM products WHERE id=?',(item['product_id'],))
                if not product:
                    raise ValueError('Produto do pedido não encontrado.')
                recipe=self.recipe(item['product_id'])
                if not recipe or any(r['material_name'] is None for r in recipe):
                    raise ValueError(f"Revise a composição do produto {product['code']} antes de registrar a retirada de estoque.")
                selections={int(k):int(v) for k,v in (item.get('variants') or {}).items()}
                for r in recipe:
                    material=self.one('SELECT id FROM materials WHERE code=? COLLATE NOCASE',(r['material_code'],))
                    material_id=material['id']
                    variant_id=selections.get(material_id)
                    options=self.variants(material_id,active_only=True)
                    if self.variants(material_id) and not options and variant_id is None:
                        raise ValueError(f"Ative ao menos uma variação de {r['material_name']} antes do pedido.")
                    if options and variant_id is None:
                        raise ValueError(f"Escolha a variação de {r['material_name']} no produto {product['code']}.")
                    if variant_id is not None and not self.one(
                            'SELECT 1 FROM material_variants WHERE id=? AND material_id=?',
                            (variant_id,material_id)):
                        raise ValueError(f"Variação inválida para {r['material_name']}.")
                    usage[(material_id,variant_id)]+=item['qty']*r['qty']
        with self.db:
            if id:
                self.db.execute("UPDATE orders SET customer=?,due_date=?,payment=?,paid_cents=?,production=?,notes=? WHERE id=?",
                                (customer,due_date,payment,paid_cents,production,notes.strip(),id))
                self.db.execute("DELETE FROM order_items WHERE order_id=?", (id,))
            else:
                number=self.next_number()
                id = self.db.execute("INSERT INTO orders(number,customer,created_date,due_date,payment,paid_cents,production,notes) VALUES(?,?,?,?,?,?,?,?)",
                                     (number,customer,date.today().isoformat(),due_date,payment,paid_cents,production,notes.strip())).lastrowid
                self._remember_number(number)
            for item in items:
                item_id=self.db.execute("INSERT INTO order_items(order_id,product_id,description,qty,unit_cents) VALUES(?,?,?,?,?)",
                                        (id,item.get('product_id'),item['description'].strip(),item['qty'],item['unit_cents'])).lastrowid
                for material_id,variant_id in (item.get('variants') or {}).items():
                    self.db.execute('INSERT INTO order_item_variants(item_id,material_id,variant_id) VALUES(?,?,?)',
                                    (item_id,int(material_id),int(variant_id)))
            if stock_change:
                if old_items:
                    for previous in self.all('''SELECT material_id,variant_id,SUM(delta) AS net FROM stock_movements
                                                WHERE order_id=? GROUP BY material_id,variant_id''',(id,)):
                        if abs(previous['net'])>1e-8:
                            self._record_movement(previous['material_id'],-previous['net'],'order_reversal',
                                                  f'Ajuste de itens do pedido {self.order(id)["number"]}',id,
                                                  previous['variant_id'])
                for (material_id,variant_id),used in usage.items():
                    self._record_movement(material_id,-used,'order_out',
                                          f'Pedido {self.order(id)["number"]}',id,variant_id)
        return id

    def _record_movement(self, material_id, delta, kind, notes, order_id=None, variant_id=None):
        self.db.execute('''INSERT INTO stock_movements(material_id,variant_id,order_id,delta,kind,notes,effective_date,created_at)
                           VALUES(?,?,?,?,?,?,?,?)''',
                        (material_id,variant_id,order_id,delta,kind,notes,date.today().isoformat(),
                         datetime.now().isoformat(timespec='seconds')))

    def stock(self):
        balances={(row['material_id'],row['variant_id']):row['balance'] for row in self.all(
            'SELECT material_id,variant_id,SUM(delta) AS balance FROM stock_movements GROUP BY material_id,variant_id')}
        result=[]
        for material in self.materials():
            data=dict(material);options=self.variants(material['id'])
            if options:
                for variant in options:
                    result.append(dict(data,variant_id=variant['id'],variant_name=variant['name'],
                        variant_active=variant['active'],balance=balances.get((material['id'],variant['id']),0)))
                if (material['id'],None) in balances:
                    result.append(dict(data,variant_id=None,variant_name='Sem variação (legado)',
                                       variant_active=1,balance=balances[(material['id'],None)]))
            else:
                result.append(dict(data,variant_id=None,variant_name='Padrão',variant_active=1,
                                   balance=balances.get((material['id'],None),0)))
        return result

    def stock_history(self, material_id, limit=200, variant_id=None):
        return self.all('''SELECT s.*,o.number AS order_number FROM stock_movements s
          LEFT JOIN orders o ON o.id=s.order_id
          WHERE s.material_id=? AND s.variant_id IS ? ORDER BY s.id DESC LIMIT ?''',
          (material_id,variant_id,limit))

    def adjust_stock(self, material_id, action, qty, notes, variant_id=None):
        """Manual entry, forced withdrawal or exact balance correction."""
        if not self.one('SELECT id FROM materials WHERE id=?',(material_id,)):
            raise ValueError('Insumo não encontrado.')
        if not math.isfinite(qty) or not notes.strip():
            raise ValueError('Informe uma quantidade válida e o motivo do ajuste.')
        if variant_id is not None and not self.one('SELECT 1 FROM material_variants WHERE id=? AND material_id=?',
                                                    (variant_id,material_id)):
            raise ValueError('A variação não pertence a este insumo.')
        if variant_id is None and self.variants(material_id) and not self.one(
                'SELECT 1 FROM stock_movements WHERE material_id=? AND variant_id IS NULL LIMIT 1',(material_id,)):
            raise ValueError('Escolha uma variação do insumo.')
        current=self.one('''SELECT COALESCE(SUM(delta),0) AS balance FROM stock_movements
                            WHERE material_id=? AND variant_id IS ?''',(material_id,variant_id))['balance']
        if action=='entry':
            if qty<=0:raise ValueError('A entrada deve ser maior que zero.')
            delta=qty
        elif action=='manual_out':
            if qty<=0:raise ValueError('A retirada deve ser maior que zero.')
            delta=-qty
        elif action=='adjustment':
            delta=qty-current
            if abs(delta)<1e-8:raise ValueError('O saldo informado já é o saldo atual.')
        else:raise ValueError('Tipo de movimentação inválido.')
        with self.db:
            self._record_movement(material_id,delta,action,notes.strip(),variant_id=variant_id)
        return current+delta

    def transfer_stock(self, material_id, source_variant_id, target_variant_id, qty, notes):
        """Move existing balance between colors without changing total stock."""
        if source_variant_id==target_variant_id or target_variant_id is None:
            raise ValueError('Escolha uma variação de destino diferente da origem.')
        if not math.isfinite(qty) or qty<=0 or not notes.strip():
            raise ValueError('Informe uma quantidade positiva e o motivo da transferência.')
        if not self.one('SELECT 1 FROM material_variants WHERE id=? AND material_id=? AND active=1',
                        (target_variant_id,material_id)):
            raise ValueError('Variação de destino inválida ou inativa.')
        if source_variant_id is not None and not self.one(
                'SELECT 1 FROM material_variants WHERE id=? AND material_id=?',
                (source_variant_id,material_id)):
            raise ValueError('Variação de origem inválida.')
        current=self.one('''SELECT COALESCE(SUM(delta),0) AS balance FROM stock_movements
                            WHERE material_id=? AND variant_id IS ?''',(material_id,source_variant_id))['balance']
        if current+1e-8<qty:
            raise ValueError('A origem não tem saldo suficiente para transferir. Faça um ajuste de estoque se necessário.')
        with self.db:
            self._record_movement(material_id,-qty,'transfer_out',notes.strip(),variant_id=source_variant_id)
            self._record_movement(material_id,qty,'transfer_in',notes.strip(),variant_id=target_variant_id)

    def mean_price(self, product_id):
        row = self.one("SELECT SUM(qty*unit_cents)*1.0/SUM(qty) FROM order_items WHERE product_id=? AND unit_cents IS NOT NULL", (product_id,))
        return round(row[0]) if row and row[0] is not None else None

    def dashboard(self):
        orders = self.orders()
        return {
            'materials': self.one('SELECT COUNT(*) FROM materials WHERE active=1')[0],
            'products': self.one('SELECT COUNT(*) FROM products WHERE active=1')[0],
            'orders': len(orders),
            'open': sum(o['production'] != 'Entregue' for o in orders),
            'revenue': sum(o['total_cents'] or 0 for o in orders if not o['unpriced']),
            'mean_item_price': self.one('SELECT ROUND(SUM(qty*unit_cents)*1.0/SUM(qty)) FROM order_items WHERE unit_cents IS NOT NULL')[0],
            'due': [o for o in orders if o['production'] != 'Entregue'][:8]
        }

    def backup(self, destination):
        other=sqlite3.connect(destination)
        try:
            self.db.backup(other)
        finally:
            other.close()

    def report_data(self, month=None, year=None):
        """Monthly sales use order creation dates; Total includes undated imports."""
        if month is not None and (not isinstance(month,int) or month not in range(1,13) or
                                  not isinstance(year,int) or year not in range(1900,10000)):
            raise ValueError('Selecione um mês e ano válidos.')
        prefix=f'{year:04d}-{month:02d}' if month is not None else None
        orders=[o for o in self.orders() if prefix is None or
                (o['created_date'] and o['created_date'].startswith(prefix))]
        priced=[o for o in orders if not o['unpriced']]
        revenue=sum(round(o['total_cents'] or 0) for o in priced)
        ids=[o['id'] for o in priced]
        products=[]
        if ids:
            placeholders=','.join('?' for _ in ids)
            products=self.all(f"""SELECT COALESCE(p.name,'Avulso: '||i.description) AS name,
              SUM(i.qty) AS quantity, SUM(ROUND(i.qty*i.unit_cents)) AS total_cents
              FROM order_items i JOIN orders o ON o.id=i.order_id
              LEFT JOIN products p ON p.id=i.product_id
              WHERE i.order_id IN ({placeholders}) AND i.unit_cents IS NOT NULL
              GROUP BY COALESCE('P'||p.id,'A'||i.description)
              ORDER BY total_cents DESC, name COLLATE NOCASE""",ids)
        def counts(key):
            result={}
            for row in orders:
                label=row[key] or 'Não informado'
                result[label]=result.get(label,0)+1
            return sorted(result.items(),key=lambda value:(-value[1],value[0]))
        return {'orders':orders,'count':len(orders),'open':sum(o['production']!='Entregue' for o in orders),
                'unknown':len(orders)-len(priced),'revenue':revenue,
                'average_ticket':round(revenue/len(priced)) if priced else None,
                'payment':counts('payment'),'production':counts('production'),'products':products,
                'month':month,'year':year}

    def import_workbook(self, path):
        """Import once, preserving unresolved recipe codes and historical orders."""
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            with self.db:
                for row in workbook['Insumos'].iter_rows(min_row=2, max_col=6, values_only=True):
                    code, name, pack_qty, size, weight, price = row
                    if not name or not pack_qty or price is None:
                        continue
                    def stringify(x):
                        return str(x).rstrip('0').rstrip('.') if isinstance(x,float) else str(x) if x is not None else ''
                    self.db.execute("INSERT OR IGNORE INTO materials(code,name,size,grammage,pack_qty,unit,pack_cents) VALUES(?,?,?,?,?,?,?)",
                                    (str(code or name).strip().upper(),str(name).strip(),stringify(size),stringify(weight),float(pack_qty),'un',round(float(price)*100)))
                for row in workbook['Precificação'].iter_rows(min_row=2, max_col=25, values_only=True):
                    if not row[0] or not row[1]:
                        continue
                    pid = self.db.execute('INSERT OR IGNORE INTO products(code,name,markup,table_cents) VALUES(?,?,?,?)',
                                          (str(row[1]).strip(),str(row[0]).strip(),1.8,round(float(row[24])*100) if isinstance(row[24],(int,float)) else None)).lastrowid
                    if not pid:
                        continue
                    pos = 0
                    for idx in range(2,22,2):
                        if row[idx] and row[idx+1]:
                            self.db.execute('INSERT INTO recipes VALUES(?,?,?,?)',
                                            (pid,pos,str(row[idx]).strip().upper(),float(row[idx+1])))
                            pos += 1
                for row in workbook['Pedidos'].iter_rows(min_row=2,max_col=8,values_only=True):
                    number, customer, desc, qty, price, _total, payment, production = row
                    if not number or not customer or not desc or not qty:
                        continue
                    existing = self.db.execute("SELECT id FROM products WHERE name=? COLLATE NOCASE", (str(desc).strip(),)).fetchone()
                    oid = self.db.execute('INSERT OR IGNORE INTO orders(number,customer,created_date,due_date,payment,production,notes) VALUES(?,?,?,?,?,?,?)',
                                           (str(number).strip(),str(customer).strip(),None,None,str(payment or 'Pendente'),
                                            str(production or 'Novo'), 'Importado da planilha; data de entrega não informada.')).lastrowid
                    if oid:
                        self.db.execute('INSERT INTO order_items(order_id,product_id,description,qty,unit_cents) VALUES(?,?,?,?,?)',
                                        (oid, existing['id'] if existing else None,str(desc).strip(),float(qty),round(float(price)*100) if price is not None else None))
        finally:
            workbook.close()


def export_order_pdf(store: Store, order_id: int, destination: str | Path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader

    order = store.order(order_id)
    if order is None:
        raise ValueError('Pedido não encontrado.')
    items = store.items(order_id)
    pdfmetrics.registerFont(TTFont('DejaVuERP', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')) if os.path.isfile('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf') else None
    font = 'DejaVuERP' if 'DejaVuERP' in pdfmetrics.getRegisteredFontNames() else 'Helvetica'
    bold = 'Helvetica-Bold'
    cv = canvas.Canvas(str(destination), pagesize=A4)
    width,height = A4
    red,olive,gold = colors.HexColor('#A7463E'),colors.HexColor('#48452C'),colors.HexColor('#BD8738')

    def header():
        cv.setFillColor(colors.HexColor('#FFF5EA'));cv.rect(0,height-125,width,125,fill=1,stroke=0)
        cv.drawImage(ImageReader(str(asset('logo.png'))),38,height-112,width=98,height=98,mask='auto')
        cv.setFillColor(olive);cv.setFont(bold,19);cv.drawString(145,height-60,'PAPÉIS DE MARTE')
        cv.setFillColor(red);cv.setFont(font,11);cv.drawString(146,height-83,'PEDIDO ' + order['number'])
        cv.setStrokeColor(gold);cv.line(38,height-130,width-38,height-130)
    def line(text,x,y,maxwidth):
        while text:
            part=text
            while pdfmetrics.stringWidth(part,font,9)>maxwidth and len(part)>1:
                part=part[:-1]
            if len(part)<len(text):
                boundary=part.rfind(' ')
                if boundary>0:part=part[:boundary]
            cv.drawString(x,y,part);text=text[len(part):].lstrip();y-=14
        return y
    def footer():
        cv.setStrokeColor(gold);cv.line(38,44,width-38,44)
        cv.setFillColor(olive);cv.setFont(font,8);cv.drawString(38,29,'Papéis de Marte • comprovante do pedido')
        cv.drawRightString(width-38,29,f'Página {cv.getPageNumber()}')

    header();y=height-155
    cv.setFillColor(olive);cv.setFont(bold,10);cv.drawString(38,y,'CLIENTE');cv.setFont(font,10)
    y=line(order['customer'],114,y,445)-17
    cv.setFont(bold,10);cv.drawString(38,y,'PEDIDO');cv.setFont(font,10)
    cv.drawString(114,y,br_date(order['created_date']))
    cv.setFont(bold,10);cv.drawString(318,y,'ENTREGA');cv.setFont(font,10);cv.drawString(390,y,br_date(order['due_date']))
    y-=28
    cv.setFont(bold,10);cv.drawString(38,y,'PAGAMENTO');cv.setFont(font,10);cv.drawString(114,y,order['payment'])
    cv.setFont(bold,10);cv.drawString(318,y,'PRODUÇÃO');cv.setFont(font,10);cv.drawString(390,y,order['production'])
    y-=36
    def table_head(y):
        cv.setFillColor(olive);cv.rect(38,y-8,width-76,25,fill=1,stroke=0)
        cv.setFillColor(colors.white);cv.setFont(bold,9)
        cv.drawString(48,y,'ITEM');cv.drawRightString(375,y,'QTD')
        cv.drawRightString(460,y,'UNIT.');cv.drawRightString(width-48,y,'TOTAL')
        cv.setFillColor(olive)
        return y-28
    y=table_head(y)
    total=0; unknown=False
    for item in items:
        variants='; '.join(store.variant_labels(item.get('variants')))
        desc=item['description']+(f'  |  {variants}' if variants else '')
        nlines=max(1, (len(desc)+37)//38)
        if y-nlines*15 < 90:
            footer();cv.showPage();header();y=table_head(height-160)
        cv.setFont(font,9)
        y=line(desc,48,y,235)
        cv.drawRightString(375,y+14*nlines,fmt_qty(item['qty']))
        cv.drawRightString(460,y+14*nlines,money(item['unit_cents']))
        line_total=round(item['qty']*item['unit_cents']) if item['unit_cents'] is not None else None
        cv.drawRightString(width-48,y+14*nlines,money(line_total))
        if line_total is None:unknown=True
        else:total+=line_total
        cv.setStrokeColor(colors.HexColor('#EAD9CC'));cv.line(38,y-3,width-38,y-3)
        y-=19
    if y<130:footer();cv.showPage();header();y=height-165
    cv.setFont(bold,13);cv.setFillColor(red)
    cv.drawRightString(width-48,y-10,'TOTAL: ' + (money(total) if not unknown else 'A definir'))
    y-=41
    if order['payment']=='Parcial' and not unknown:
        cv.setFont(font,10);cv.setFillColor(olive)
        cv.drawRightString(width-48,y,'RECEBIDO: '+money(order['paid_cents']))
        y-=18
        cv.setFont(bold,10);cv.setFillColor(red)
        cv.drawRightString(width-48,y,'RESTANTE: '+money(max(0,total-order['paid_cents'])))
        y-=28
    if order['notes']:
        cv.setFillColor(olive);cv.setFont(bold,10);cv.drawString(38,y,'OBSERVAÇÕES');y-=17
        cv.setFont(font,9)
        for paragraph in order['notes'].splitlines():
            if y<85:footer();cv.showPage();header();y=height-165
            y=line(paragraph,38,y,width-76)-5
    footer();cv.save()


def export_report_pdf(data: dict, destination: str | Path):
    """Produce a paginated management report using the same order totals as the UI."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.utils import ImageReader
    from xml.sax.saxutils import escape

    red,olive,gold = (colors.HexColor(v) for v in ('#A7463E','#48452C','#BD8738'))
    light=colors.HexColor('#FFF5EA')
    text_style=ParagraphStyle('cell',fontName='Helvetica',fontSize=8,leading=11,textColor=olive)
    head_style=ParagraphStyle('head',parent=text_style,fontName='Helvetica-Bold',textColor=colors.white)
    right_style=ParagraphStyle('right',parent=text_style,alignment=TA_RIGHT)
    center_style=ParagraphStyle('center',parent=text_style,alignment=TA_CENTER)
    title_style=ParagraphStyle('title',fontName='Helvetica-Bold',fontSize=13,leading=18,textColor=red,spaceAfter=8)
    subtle_style=ParagraphStyle('subtle',fontName='Helvetica',fontSize=8,leading=12,textColor=olive)
    pdf=SimpleDocTemplate(str(destination),pagesize=A4,leftMargin=40,rightMargin=40,
                          topMargin=123,bottomMargin=53,title='Relatório - Papéis de Marte')
    available=A4[0]-80
    story=[]

    def header(canvas,doc):
        canvas.saveState()
        w,h=A4
        canvas.setFillColor(light);canvas.rect(0,h-102,w,102,fill=1,stroke=0)
        canvas.drawImage(ImageReader(str(asset('logo.png'))),38,h-93,width=83,height=83,mask='auto')
        canvas.setFillColor(olive);canvas.setFont('Helvetica-Bold',18)
        canvas.drawString(135,h-46,'PAPÉIS DE MARTE')
        canvas.setFillColor(red);canvas.setFont('Helvetica',10)
        canvas.drawString(136,h-67,'RELATÓRIO DE PEDIDOS')
        canvas.setStrokeColor(gold);canvas.line(40,h-105,w-40,h-105)
        canvas.line(40,43,w-40,43)
        canvas.setFillColor(olive);canvas.setFont('Helvetica',8)
        canvas.drawString(40,29,'Dados registrados no aplicativo local')
        canvas.drawRightString(w-40,29,f'Página {doc.page}')
        canvas.restoreState()

    def p(value,style=text_style):
        return Paragraph(escape(str(value)),style)

    def table(rows,widths,header_row=True):
        t=Table(rows,colWidths=widths,repeatRows=1 if header_row else 0,hAlign='LEFT')
        commands=[('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),
                  ('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),
                  ('BOTTOMPADDING',(0,0),(-1,-1),7),('LINEBELOW',(0,-1),(-1,-1),0.5,gold)]
        if header_row:
            commands.extend([('BACKGROUND',(0,0),(-1,0),olive),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,light])])
        else:
            commands.append(('BACKGROUND',(0,0),(-1,-1),light))
        t.setStyle(TableStyle(commands))
        return t

    months=('Janeiro','Fevereiro','Março','Abril','Maio','Junho','Julho','Agosto',
            'Setembro','Outubro','Novembro','Dezembro')
    period=(f'{months[data["month"]-1]} de {data["year"]}' if data['month'] else 'Total • todo o histórico')
    story.append(Paragraph('Resumo do período',title_style))
    story.append(Paragraph('Período por data de cadastro: '+escape(period),subtle_style))
    story.append(Spacer(1,14))
    kpis=[['Pedidos','Em aberto','Total com preço','Ticket médio'],
          [str(data['count']),str(data['open']),money(data['revenue']),money(data['average_ticket'])]]
    summary=table([[p(v,center_style) for v in row] for row in kpis],[available/4]*4,False)
    summary.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),light),
         ('TEXTCOLOR',(0,1),(-1,1),red),('TOPPADDING',(0,0),(-1,-1),10),('BOTTOMPADDING',(0,0),(-1,-1),10)]))
    story.extend([summary,Spacer(1,10)])
    if data['unknown']:
        story.append(Paragraph(f"Pedidos com valor a definir: {data['unknown']}. Excluídos dos valores totais, do ticket médio e dos itens vendidos.",subtle_style))
        story.append(Spacer(1,9))
    if data['month']:
        story.append(Paragraph('Pedidos históricos sem data de cadastro aparecem apenas na opção Total.',subtle_style))
        story.append(Spacer(1,8))

    def section(title,headers,records,widths):
        story.append(Spacer(1,13));story.append(Paragraph(title,title_style))
        if records:
            story.append(table([[p(h,head_style) for h in headers]]+
                               [[p(value,right_style if j==len(row)-1 else text_style) for j,value in enumerate(row)]
                                for row in records],widths))
        else:
            story.append(Paragraph('Nenhum registro neste período.',subtle_style))

    section('Situação dos pedidos',['Pagamento','Pedidos'],data['payment'],[available-100,100])
    section('Etapas da produção',['Situação','Pedidos'],data['production'],[available-100,100])
    section('Itens vendidos',['Produto ou item avulso','Quantidade','Valor registrado'],
            [(r['name'],fmt_qty(r['quantity']),money(round(r['total_cents']))) for r in data['products']],
            [available-193,83,110])
    section('Pedidos',['Número','Cliente','Entrega','Pagamento','Produção','Total'],
            [(r['number'],r['customer'],br_date(r['due_date']),r['payment'],r['production'],
              money(round(r['total_cents'])) if not r['unpriced'] else 'A definir') for r in data['orders']],
            [55,115,70,80,105,available-425])
    pdf.build(story,onFirstPage=header,onLaterPages=header)
